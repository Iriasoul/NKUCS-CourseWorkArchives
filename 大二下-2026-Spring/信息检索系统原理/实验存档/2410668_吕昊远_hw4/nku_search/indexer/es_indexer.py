"""
Elasticsearch 索引构建器
=======================
- 读取 crawler 产出的 pages.jsonl / docs.jsonl
- 对 PDF / DOCX / XLSX 抽取文本
- 批量写入 ES
- 把 PageRank 分数 (data/pagerank.json) 写回每条页面文档
"""
from __future__ import annotations

import io
import json
import os
import sys
from typing import Iterator

from elasticsearch import Elasticsearch, helpers

sys.path.append(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
import config
from indexer.mapping import PAGE_MAPPING, DOC_MAPPING, LOG_MAPPING

# ES 客户端
def get_es() -> Elasticsearch:
    kwargs = {"hosts": [config.ES_HOST], "request_timeout": 60}
    if config.ES_USERNAME and config.ES_PASSWORD:
        kwargs["basic_auth"] = (config.ES_USERNAME, config.ES_PASSWORD)
    return Elasticsearch(**kwargs)


# 索引创建

def ensure_index(es: Elasticsearch, name: str, body: dict, recreate: bool = False):
    exists = es.indices.exists(index=name)
    if exists and recreate:
        es.indices.delete(index=name)
        exists = False
    if not exists:
        es.indices.create(index=name, body=body)
        print(f"[index] created {name}")


# 文档文本抽取
def extract_pdf(path: str) -> str:
    try:
        from pdfminer.high_level import extract_text
        return extract_text(path) or ""
    except Exception as e:
        return ""


def extract_docx(path: str) -> str:
    try:
        from docx import Document
        doc = Document(path)
        return "\n".join(p.text for p in doc.paragraphs)
    except Exception:
        return ""


def extract_xlsx(path: str) -> str:
    try:
        from openpyxl import load_workbook
        wb = load_workbook(path, read_only=True, data_only=True)
        parts = []
        for ws in wb.worksheets:
            for row in ws.iter_rows(values_only=True):
                parts.append(" ".join(str(c) for c in row if c is not None))
        return "\n".join(parts)
    except Exception:
        return ""


def extract_text(local_path: str, ext: str) -> str:
    ext = ext.lower().lstrip(".")
    if ext == "pdf":
        return extract_pdf(local_path)[:200_000]
    if ext == "docx":
        return extract_docx(local_path)[:200_000]
    if ext == "xlsx":
        return extract_xlsx(local_path)[:200_000]
    return ""


# 批量索引
def iter_pages(path: str, pagerank: dict | None = None) -> Iterator[dict]:
    with open(path, "r", encoding="utf-8") as f:
        for line in f:
            line = line.strip()
            if not line:
                continue
            rec = json.loads(line)
            url = rec["url"]
            doc = {
                "_index": config.INDEX_PAGES,
                "_id": rec["doc_id"],
                "_source": {
                    "doc_id": rec["doc_id"],
                    "url": url,
                    "site": rec.get("site", ""),
                    "title": rec.get("title", ""),
                    "body": rec.get("body", ""),
                    "anchor_text": rec.get("anchor_text", ""),
                    "out_links": rec.get("out_links", []),
                    "snapshot_path": rec.get("snapshot_path", ""),
                    "crawl_time": int(rec.get("crawl_time", 0)),
                    "publish_time": int(rec.get("publish_time", 0)),
                    "pagerank": float(pagerank.get(url, 0.0)) if pagerank else 0.0,
                }
            }
            yield doc


def iter_docs(path: str) -> Iterator[dict]:
    with open(path, "r", encoding="utf-8") as f:
        for line in f:
            line = line.strip()
            if not line:
                continue
            rec = json.loads(line)
            content = extract_text(rec["local_path"], rec["ext"])
            yield {
                "_index": config.INDEX_DOCS,
                "_id": rec["doc_id"],
                "_source": {
                    "doc_id": rec["doc_id"],
                    "url": rec["url"],
                    "site": rec.get("site", ""),
                    "title": rec.get("title", ""),
                    "content": content,
                    "ext": rec.get("ext", ""),
                    "local_path": rec.get("local_path", ""),
                    "referer": rec.get("referer", ""),
                    "crawl_time": int(rec.get("crawl_time", 0)),
                }
            }


def bulk_index_pages(es: Elasticsearch, jsonl_path: str,
                     pagerank: dict | None = None, chunk_size: int = 500):
    success, failed = 0, 0
    for ok, result in helpers.streaming_bulk(
        es, iter_pages(jsonl_path, pagerank), chunk_size=chunk_size,
        raise_on_error=False
    ):
        if ok:
            success += 1
        else:
            failed += 1
        if (success + failed) % 2000 == 0:
            print(f"[index pages] ok={success} fail={failed}")
    print(f"[index pages] DONE ok={success} fail={failed}")


def bulk_index_docs(es: Elasticsearch, jsonl_path: str, chunk_size: int = 50):
    success, failed = 0, 0
    for ok, result in helpers.streaming_bulk(
        es, iter_docs(jsonl_path), chunk_size=chunk_size, raise_on_error=False
    ):
        if ok:
            success += 1
        else:
            failed += 1
    print(f"[index docs] DONE ok={success} fail={failed}")


# 入口主函数
def main(recreate: bool = False):
    es = get_es()
    ensure_index(es, config.INDEX_PAGES, PAGE_MAPPING, recreate=recreate)
    ensure_index(es, config.INDEX_DOCS, DOC_MAPPING, recreate=recreate)
    ensure_index(es, config.INDEX_LOGS, LOG_MAPPING, recreate=recreate)
    # PageRank
    pagerank = {}
    if os.path.exists(config.PAGERANK_PATH):
        with open(config.PAGERANK_PATH, "r", encoding="utf-8") as f:
            pagerank = json.load(f)
        print(f"[index] loaded pagerank for {len(pagerank)} urls")

    # 网页
    pages_path = os.path.join(config.DATA_DIR, "pages.jsonl")
    if os.path.exists(pages_path):
        bulk_index_pages(es, pages_path, pagerank)

    # 文档
    docs_path = os.path.join(config.DATA_DIR, "docs.jsonl")
    if os.path.exists(docs_path):
        bulk_index_docs(es, docs_path)


if __name__ == "__main__":
    import argparse
    p = argparse.ArgumentParser()
    p.add_argument("--recreate", action="store_true",
                   help="Drop and rebuild indices.")
    args = p.parse_args()
    main(recreate=args.recreate)
